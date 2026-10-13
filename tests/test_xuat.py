"""Xuất bản hợp nhất ra Markdown/Word (src/query/xuat.py, scripts/xuat_bo_luat.py) trên bộ luật thử nghiệm."""
import json

import pytest

from src.apply import hieu_dinh as hd
from src.apply.rebuild import rebuild
from src.query.lookup import get_cay
from src.query.xuat import _doan, dung_bo_luat, khoi_noi_dung, ra_docx, ra_markdown

from test_apply_rebuild import D1, D2, D3, P, env  # noqa: F401  (env là fixture)
from test_cli_vbhn import dung, run

L12 = dict(nguon="12/2017/QH14", ban_hanh="2017-06-20")
L86 = dict(nguon="86/2025/QH15", ban_hanh="2025-06-25")


@pytest.fixture
def daapdung(env):  # noqa: F811
    env.cmd("sua_doi", {"muc_tieu": [{"duong_dan": P("dieu", "2", "khoan", "2")}]}, tu=D1,
            noi_dung="2. Người phạm tội tự thú được miễn trách nhiệm hình sự.", vi_tri="Điều 1 khoản 3", **L12)
    env.cmd("bo_sung", {"cha": P("chuong", "II"), "sau": P("dieu", "3"), "loai": "dieu", "ky_hieu": "3a"}, tu=D2,
            noi_dung="Điều 3a. Tội mới\n1. Người nào phạm tội mới:\na) bị phạt tiền;\nb) bị phạt tù.", vi_tri="Điều 1", **L86)
    env.cmd("bai_bo", {"muc_tieu": [P("dieu", "4")]}, tu=D3, vi_tri="Điều 2 khoản 1", **L86)
    hd.luu_de_xuat(env.c, env.vb, [
        {"khoa": "dieu 3/khoan 1", "truong": "noi_dung", "tu_cu": "trộm cắp tài sản", "thanh_moi": "trộm cắp tài sản của người khác", "loai": "nhieu_chu",
         "bang_chung": {"nguon_so_sanh": "135/VBHN-VPQH (bản hệ thống ngày 2026-01-01)", "tham_chieu": {"11/VBHN-VPQH": "cung_sua"}}},
        {"khoa": "dieu 1", "truong": "noi_dung", "tu_cu": "tội phạm", "thanh_moi": "các tội phạm", "loai": "loi_go", "bang_chung": {}}])
    hd.duyet(env.c, env.vb, nguoi="Người duyệt", ids=[1])
    env.c.commit()
    assert rebuild(env.c, env.vb).n_ok == 3
    return env


def test_cay_theo_thu_tu_va_theo_ngay(daapdung):
    env = daapdung
    roots = get_cay(env.c, env.vb, "2026-10-02")
    assert [r.loai for r in roots] == ["chuong", "chuong"] and [c.ky_hieu for c in roots[1].children] == ["3", "3a", "5"]          # Điều 4 đã bãi bỏ từ D3 (2026-01-01)
    assert [c.ky_hieu for c in get_cay(env.c, env.vb, "2025-12-31")[1].children] == ["3", "3a", "4", "5"]
    assert [c.ky_hieu for c in get_cay(env.c, env.vb, "2017-12-31")[1].children] == ["3", "4", "5"]                               # 3a chưa có
    assert get_cay(env.c, env.vb, "2016-06-30") == []                                                                         # trước ngày hiệu lực của văn bản


def test_bo_luat_thong_ke_va_nguon(daapdung):
    bl = dung_bo_luat(daapdung.c, daapdung.vb, "2026-10-02")
    assert (bl.n_dieu, bl.n_dieu_da_sua) == (5, 2)                                                           # Điều 1,2,3,3a,5: 2 và 3a có nguồn
    assert [(v["so_hieu"], v["n_lenh"], v["tu"], v["den"]) for v in bl.cac_van_ban] == [("12/2017/QH14", 1, D1, D1), ("86/2025/QH15", 2, D2, D3)]
    assert [(r["id"], r["dieu"], r["nhan"]) for r in bl.hieu_dinh] == [(1, "3", "khoản 1")]
    assert (bl.lenh_cho_duyet, bl.hieu_dinh_cho_duyet) == (0, 1) and bl.luu_y and "01/07/2016" in bl.luu_y
    truoc = dung_bo_luat(daapdung.c, daapdung.vb, "2017-12-31")                                              # trước 12/2017: chỉ văn bản gốc
    assert (truoc.n_dieu, truoc.n_dieu_da_sua, truoc.cac_van_ban) == (5, 0, []) and len(truoc.hieu_dinh) == 1 and truoc.hieu_dinh[0]["dieu"] == "3"
    with pytest.raises(ValueError):
        dung_bo_luat(daapdung.c, 9999, "2026-10-02")


def test_markdown(daapdung):
    dc = {"nhan": "135/VBHN-VPQH", "ngay": "2026-10-02", "n_theirs": 20, "n_match": 18, "n_match_fold": 19, "n_noi_dung": 1, "n_trinh_bay": 1, "dieu_lech": ["2"]}
    md = ra_markdown(dung_bo_luat(daapdung.c, daapdung.vb, "2026-10-02", dc))
    assert md.startswith("# ") and "BẢN HỢP NHẤT TẠI NGÀY 02/10/2026" in md and "KHÔNG phải văn bản hợp nhất chính thức" in md
    assert "| Luật số 12/2017/QH14" not in md and "(số 12/2017/QH14) | 20/06/2017 | 01/01/2018 | 1 |" in md          # bảng cơ sở hợp nhất
    assert "19/20 nút khớp (95.00%)" in md and "18 nút giống hệt từng chữ" in md and "Còn 1 chỗ khác về nội dung ở các Điều: 2" in md
    assert "Sự khớp này đạt được SAU 1 hiệu đính kỹ thuật" in md
    assert "###### Điều 2. Nguyên tắc" in md and "2\\. Người phạm tội tự thú được miễn trách nhiệm hình sự." in md
    assert "&emsp;&emsp;a) phạt cảnh cáo;" in md
    assert "> *Nguồn sửa đổi — khoản 2: sửa đổi, bổ sung bởi" in md and "Điều 1 khoản 3), hiệu lực từ 01/01/2018*" in md
    assert "> *Nguồn sửa đổi — toàn Điều: bổ sung bởi" in md                                                  # Điều 3a mới
    assert "###### Điều 4." not in md                                                                          # đã bãi bỏ
    assert "> *Hiệu đính kỹ thuật — khoản 1: “trộm cắp tài sản” → “trộm cắp tài sản của người khác” (theo 135/VBHN-VPQH; duyệt: Người duyệt)*" in md
    assert "Trộm" not in md.split("## Phụ lục")[0].split("Điều 3.")[0]
    phu_luc = md.split("## Phụ lục")[1]
    assert "| 1 | Điều 3 — khoản 1 | “trộm cắp tài sản” | “trộm cắp tài sản của người khác” | nhiều chữ | 135/VBHN-VPQH; đồng thuận: 11/VBHN-VPQH | Người duyệt" in phu_luc
    assert "Còn 1 đề xuất hiệu đính kỹ thuật CHƯA được duyệt" in md and "chưa nạp nghị quyết hoãn/lùi hiệu lực nào" in md
    assert md.count("|---|") >= 2 and "\n\n\n" not in md


def test_markdown_khong_doi_chieu_va_khong_hieu_dinh(env):  # noqa: F811
    md = ra_markdown(dung_bo_luat(env.c, env.vb, "2026-10-02"))
    assert "Chưa chạy đối chiếu" in md and "Không có hiệu đính kỹ thuật nào" in md and "Nguồn sửa đổi" not in md
    assert "Quy mô: 5 Điều" in md and "0 Điều đã được sửa đổi" in md


def test_doan_tach_nhieu_doan():
    assert _doan("a\nb\n\n c ", 1, 0) == [("doan", 1, "a"), ("doan", 0, "b"), ("doan", 0, "c")]
    assert _doan("", 0, 0) == []


def test_khoi_noi_dung_thu_tu_va_dung_loai(daapdung):
    kinds = [k for k, _, _ in khoi_noi_dung(dung_bo_luat(daapdung.c, daapdung.vb, "2026-10-02"))]
    assert kinds[0] == "tieu_de_van_ban" and kinds.count("tieu_de_van_ban") == 1 and kinds[-1] == "bang"


def test_docx(daapdung, tmp_path):
    import docx

    out = ra_docx(dung_bo_luat(daapdung.c, daapdung.vb, "2026-10-02"), tmp_path / "x" / "bo.docx")
    d = docx.Document(out)
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Điều 2. Nguyên tắc" in heads and "Điều 3a. Tội mới" in heads and "Chương II — CÁC TỘI DANH" in heads and "Điều 4. Tội lừa đảo" not in heads
    assert len(d.tables) == 2 and d.tables[0].cell(0, 0).text == "Văn bản" and d.tables[1].cell(1, 2).text == "“trộm cắp tài sản”"
    assert any(p.text.startswith("Nguồn sửa đổi — khoản 2") and p.runs[0].italic for p in d.paragraphs)


def test_cli_xuat(tmp_path):
    raw, db, _ = dung(tmp_path)
    ra = tmp_path / "xuat"
    out = run("xuat_bo_luat.py", "--dich", "99/2098/QH99", "--ngay", "2026-10-02", "--ra", ra, "--db", db, raw=raw)
    assert "5 Điều có hiệu lực tại 2026-10-02" in out and "chưa đối chiếu" in out
    stem = next(f for f in ra.glob("*.md")).stem
    assert stem.endswith("_2026-10-02") and "99-2098-QH99" in stem and (ra / (stem + ".docx")).exists()
    # có văn bản hợp nhất giả lập -> bản xuất ghi kết quả đối chiếu
    from test_cli_vbhn import THAN, ghi_vbhn
    ghi_vbhn(raw, "135-VBHN-VPQH", THAN)
    out = run("xuat_bo_luat.py", "--dich", "99/2098/QH99", "--ngay", "2026-10-02", "--ra", ra, "--dinh-dang", "md", "--db", db, raw=raw)
    assert "đối chiếu 135/VBHN-VPQH: 17/17 nút khớp, 0 chỗ khác nội dung" in out
    assert "17/17 nút khớp (100.00%)" in (ra / (stem + ".md")).read_text(encoding="utf-8")
    run("xuat_bo_luat.py", "--dich", "99/2098/QH99", "--ngay", "2000-01-01", "--ra", ra, "--db", db, raw=raw, expect=1)        # trước ngày hiệu lực
    run("xuat_bo_luat.py", "--dich", "1/2000/XX", "--db", db, raw=raw, expect=2)
    run("xuat_bo_luat.py", "--db", tmp_path / "khong.sqlite", raw=raw, expect=2)


def test_xem_truoc_gom_ca_lenh_va_hieu_dinh_chua_duyet(env):  # noqa: F811
    from src.apply.rebuild import preview

    env.cmd("sua_doi", {"muc_tieu": [{"duong_dan": P("dieu", "2", "khoan", "2")}]}, tu=D1, trang_thai="cho_duyet",
            noi_dung="2. Người phạm tội tự thú được miễn trách nhiệm hình sự.", vi_tri="Điều 1 khoản 3", **L12)
    hd.luu_de_xuat(env.c, env.vb, [{"khoa": "dieu 1", "truong": "noi_dung", "tu_cu": "tội phạm", "thanh_moi": "các tội phạm", "loai": "loi_go", "bang_chung": {}}])
    env.c.commit()
    # bản chính thức: chưa duyệt nên không có gì
    chinh = ra_markdown(dung_bo_luat(env.c, env.vb, "2026-10-02"))
    assert "miễn trách nhiệm" not in chinh and "XEM TRƯỚC" not in chinh and "Còn 1 lệnh sửa đổi đã nạp nhưng CHƯA được duyệt" in chinh
    # bản xem trước: có, và ghi rõ chưa duyệt
    rep, mem = preview(env.c, env.vb)
    try:
        bl = dung_bo_luat(mem, env.vb, "2026-10-02", xem_truoc=True)
    finally:
        mem.close()
    md = ra_markdown(bl)
    assert rep.n_ok == 1 and bl.xem_truoc and (bl.lenh_cho_duyet, bl.hieu_dinh_cho_duyet) == (1, 1)
    assert "(BẢN XEM TRƯỚC, CHƯA DUYỆT)" in md.splitlines()[0] and "BẢN XEM TRƯỚC — CHƯA ĐƯỢC NGƯỜI DUYỆT" in md and "đã gồm 1 lệnh sửa đổi và 1 quy tắc hiệu đính" in md
    assert "2\\. Người phạm tội tự thú được miễn trách nhiệm hình sự." in md and "(số 12/2017/QH14) | 20/06/2017 | 01/01/2018 | 1 |" in md
    assert "(theo văn bản hợp nhất; CHƯA DUYỆT)" in md and md.count("CHƯA DUYỆT") >= 2                       # ghi chú cạnh Điều và dòng phụ lục
    assert "nên chưa nằm trong bản này" not in md


def test_cli_xuat_xem_truoc(tmp_path):
    raw, db, _ = dung(tmp_path)
    hd_json = {"van_ban": "99-2098-QH99", "quy_tac": [{"khoa": "dieu 1", "truong": "noi_dung", "tu_cu": "tội phạm", "thanh_moi": "các tội phạm", "loai": "loi_go",
                                                         "trang_thai_duyet": "cho_duyet", "nguoi_duyet": None, "ngay_duyet": None, "ghi_chu_duyet": None, "bang_chung": {}}]}
    (raw / "99-2098-QH99" / "hieu_dinh.json").write_text(json.dumps(hd_json, ensure_ascii=False), encoding="utf-8")
    run("doi_chieu_vbhn.py", "nap", "--dich", "99/2098/QH99", "--db", db, raw=raw)
    ra = tmp_path / "xuat"
    out = run("xuat_bo_luat.py", "--dich", "99/2098/QH99", "--ngay", "2026-10-02", "--xem-truoc", "--dinh-dang", "md", "--ra", ra, "--db", db, raw=raw)
    f = next(ra.glob("*_XEM-TRUOC.md"))
    md = f.read_text(encoding="utf-8")
    assert "BẢN XEM TRƯỚC: đã gồm 0 lệnh sửa đổi và 1 quy tắc hiệu đính CHƯA được người duyệt" in out
    assert "các tội phạm" in md and "CHƯA DUYỆT" in md
    out = run("xuat_bo_luat.py", "--dich", "99/2098/QH99", "--ngay", "2026-10-02", "--dinh-dang", "md", "--ra", ra, "--db", db, raw=raw)      # bản chính thức: chưa duyệt -> không áp
    assert "còn 0 lệnh sửa đổi và 1 đề xuất hiệu đính chưa duyệt" in out
    assert "các tội phạm" not in next(p for p in ra.glob("*.md") if "XEM" not in p.name).read_text(encoding="utf-8").split("## Phụ lục")[0].split("## Nội dung")[1]
