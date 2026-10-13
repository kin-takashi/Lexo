"""Stage -> duyệt -> nạp -> tra cứu, trên file .docx tổng hợp; cộng thêm kiểm tra trên file BLHS thật nếu có trong inbox/."""
import json
import random
import re
import shutil
from pathlib import Path

import pytest

from src.ingest.base_law import IngestError, ingest_base_law, stage_base_law
from src.ingest.docx_text import extract_docx
from src.parse.model import Node
from src.parse.text import normalize_text
from src.query.lookup import find_van_ban, get_dieu, parse_reference, render
from src.review.base_law_review import (
    approve, blocking_problems, flag_key, load_parsed, new_review, save_review,
)
from src.store import connect, init_db
from src.review.base_law_view import dieu_spans, node_to_text, sample_dieu
from src.validate.base_law import CANH_BAO, LOI, validate_base_law

ROOT = Path(__file__).resolve().parent.parent
REAL_DOCX = ROOT / "inbox" / "Bộ-luật-100-2015-QH13.docx"

LINES = [
    "BỘ LUẬT", "THỬ NGHIỆM", "Căn cứ Hiến pháp;", "Quốc hội ban hành Bộ luật thử nghiệm.",
    "Chương I", "ĐIỀU KHOẢN CHUNG",
    "Ðiều 1. Phạm vi",                       # Ð (Latin Eth) như trong file thật
    "Bộ luật này quy định về tội phạm.",
    "Điều 2. Nguyên tắc",
    "1. Nguyên tắc như sau:", "a) tuân thủ pháp luật;", "b) bảo vệ quyền con người.",
    "2. Mức phạt tối thiểu 2.000.000 đồng.",
    "Điều 2a. Điều chèn", "Nội dung 2a.",
    "Chương II", "TIÊU ĐỀ\nHAI DÒNG",
    "Điều 3. Hiệu lực thi hành", "Bộ luật này có hiệu lực thi hành từ ngày 01 tháng 7 năm 2099.",
    "________________",
    "Bộ luật này đã được Quốc hội khóa XIII thông qua ngày 27 tháng 11 năm 2098.",
]


def build_docx(path: Path, lines, so_hieu="99/2098/QH99"):
    import docx
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "QUỐC HỘI\n______"
    t.cell(0, 1).text = f"Số: {so_hieu}"
    for line in lines:
        d.add_paragraph(line)
    d.add_paragraph("")          # đoạn rỗng phải bị bỏ
    d.save(path)
    return path


@pytest.fixture
def staged(tmp_path):
    src = build_docx(tmp_path / "BLTN.docx", LINES)
    return stage_base_law(src, tmp_path / "raw", ten_viet_tat="BLTN")


def approved_review(staged, nguoi="Người duyệt thử", mien=None, xem_canh_bao=True):
    review = new_review(staged.parsed_sha256)
    review["mien_loi"] = mien or {}
    if xem_canh_bao:
        review["da_xem_canh_bao"] = [flag_key(f) for f in staged.flags if f.level == CANH_BAO]
    save_review(approve(review, nguoi), staged.directory / "review.json")


@pytest.fixture
def conn():
    c = connect()
    init_db(c)
    yield c
    c.close()


# ---------- docx -> text ----------

def test_extract_docx(tmp_path):
    ex = extract_docx(build_docx(tmp_path / "a.docx", LINES))
    assert ex.homoglyphs_fixed == 1 and ex.n_tables == 1
    assert "99/2098/QH99" in ex.header_text
    assert "Điều 1. Phạm vi" in ex.lines and "TIÊU ĐỀ HAI DÒNG" in ex.lines     # Ð đã sửa, xuống dòng mềm đã gộp
    assert "" not in ex.lines and len(ex.sha256) == 64


def test_stage_writes_raw_copy_and_deterministic_parsed_json(tmp_path):
    src = build_docx(tmp_path / "BLTN.docx", LINES)
    a = stage_base_law(src, tmp_path / "raw", ten_viet_tat="BLTN")
    b = stage_base_law(src, tmp_path / "raw", ten_viet_tat="BLTN")
    assert a.directory.name == "99-2098-QH99" and a.raw_copy.read_bytes() == src.read_bytes()
    assert a.parsed_sha256 == b.parsed_sha256
    assert not [f for f in a.flags if f.level == LOI]


def test_stage_refuses_to_overwrite_different_raw_file(tmp_path):
    stage_base_law(build_docx(tmp_path / "BLTN.docx", LINES), tmp_path / "raw", ten_viet_tat="BLTN")
    other = build_docx(tmp_path / "other" / "BLTN.docx", LINES + ["Điều 4. Thêm"]) if (tmp_path / "other").mkdir() is None else None
    with pytest.raises(IngestError, match="KHÁC"):
        stage_base_law(other, tmp_path / "raw", ten_viet_tat="BLTN")


# ---------- điều kiện nạp ----------

def test_ingest_requires_review(staged, conn):
    with pytest.raises(IngestError, match="review.json"):
        ingest_base_law(conn, staged.directory)
    assert conn.execute("SELECT COUNT(*) FROM van_ban").fetchone()[0] == 0


def test_ingest_rejects_unapproved_or_nameless_review(staged, conn):
    save_review(new_review(staged.parsed_sha256), staged.directory / "review.json")
    with pytest.raises(IngestError, match="chưa bấm 'Duyệt'"):
        ingest_base_law(conn, staged.directory)
    with pytest.raises(ValueError):
        approve(new_review(staged.parsed_sha256), "  ")


def test_review_becomes_stale_when_parsed_json_changes(staged, tmp_path, conn):
    approved_review(staged)
    p = staged.parsed_path
    data = json.loads(p.read_text(encoding="utf-8"))
    data["preamble"].append("dòng đã sửa sau khi duyệt")
    p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(IngestError, match="bản parse cũ"):
        ingest_base_law(conn, staged.directory)


def test_ingest_detects_tampered_raw_file(staged, conn):
    approved_review(staged)
    staged.raw_copy.write_bytes(staged.raw_copy.read_bytes() + b"x")
    with pytest.raises(IngestError, match="sha256"):
        ingest_base_law(conn, staged.directory)


def test_errors_must_be_waived_with_note(tmp_path, conn):
    lines = [x for x in LINES if x != "b) bảo vệ quyền con người."] + []
    lines.insert(lines.index("a) tuân thủ pháp luật;") + 1, "c) sai thứ tự;")       # a) rồi c): lỗi
    st = stage_base_law(build_docx(tmp_path / "BLTN.docx", lines), tmp_path / "raw", ten_viet_tat="BLTN")
    errors = [f for f in st.flags if f.level == LOI]
    assert len(errors) == 1
    approved_review(st)                                                              # duyệt nhưng CHƯA miễn lỗi
    with pytest.raises(IngestError, match="chưa được miễn"):
        ingest_base_law(conn, st.directory)
    approved_review(st, mien={flag_key(errors[0]): "  "})                            # ghi chú rỗng không tính
    with pytest.raises(IngestError, match="chưa được miễn"):
        ingest_base_law(conn, st.directory)
    approved_review(st, mien={flag_key(errors[0]): "Lỗi do văn bản nguồn, đã đối chiếu bản chính thức"})
    stats = ingest_base_law(conn, st.directory)
    assert stats["so_loi_mien"] == 1
    note = conn.execute("SELECT ghi_chu FROM log_nap").fetchone()[0]
    assert "Lỗi do văn bản nguồn" in note


def test_unseen_warning_blocks_ingest(tmp_path, conn):
    lines = list(LINES)
    lines.insert(lines.index("Điều 2a. Điều chèn") + 2, "Đoạn lạc ngay dưới Chương?")   # nằm trong Điều 2a: không cảnh báo
    lines.insert(lines.index("Chương II") + 2, "1. Khoản lạc ngoài Điều")             # dưới Chương II, ngoài Điều: cảnh báo
    st = stage_base_law(build_docx(tmp_path / "BLTN.docx", lines), tmp_path / "raw", ten_viet_tat="BLTN")
    assert [f for f in st.flags if f.level == CANH_BAO]
    approved_review(st, xem_canh_bao=False)
    with pytest.raises(IngestError, match="Cảnh báo chưa được xem"):
        ingest_base_law(conn, st.directory)
    approved_review(st)
    assert ingest_base_law(conn, st.directory)["so_nut"] > 0


# ---------- nội dung DB ----------

@pytest.fixture
def loaded(staged, conn):
    approved_review(staged, nguoi="Tester")
    stats = ingest_base_law(conn, staged.directory, raw_root=staged.directory.parent)
    return conn, stats, staged


def test_ingest_populates_tables(loaded):
    conn, stats, staged = loaded
    vb = conn.execute("SELECT * FROM van_ban").fetchone()
    assert (vb["so_hieu"], vb["ten_viet_tat"]) == ("99/2098/QH99", "BLTN")
    assert (vb["ngay_ban_hanh"], vb["ngay_hieu_luc_theo_van_ban"], vb["ngay_hieu_luc"]) == ("2098-11-27", "2099-07-01", "2099-07-01")
    assert vb["file_goc"] == "99-2098-QH99/BLTN.docx"
    n_nodes = sum(1 for _ in load_parsed(staged.parsed_path).iter_nodes())
    assert stats["so_nut"] == n_nodes
    assert conn.execute("SELECT COUNT(*) FROM don_vi_luat").fetchone()[0] == n_nodes
    assert conn.execute("SELECT COUNT(*) FROM phien_ban_noi_dung").fetchone()[0] == n_nodes
    log = conn.execute("SELECT * FROM log_nap").fetchone()
    assert log["mo_hinh_su_dung"] is None and log["ket_qua"] == "thanh_cong"
    assert (log["vai_tro"], log["phuong_phap"]) == ("goc", "regex")
    pb = conn.execute("SELECT * FROM phien_ban_noi_dung").fetchall()
    assert all(r["hieu_luc_tu"] == "2099-07-01" and r["hieu_luc_den"] is None for r in pb)
    assert all(r["van_ban_tao_ra_id"] == vb["id"] and r["lenh_tao_ra_id"] is None for r in pb)
    assert all(r["nguoi_duyet"] == "Tester" and r["doi_chieu_hop_nhat"] == "chua_doi_chieu" for r in pb)


def test_ingest_twice_is_rejected(loaded):
    conn, _, staged = loaded
    with pytest.raises(IngestError, match="đã có"):
        ingest_base_law(conn, staged.directory)
    assert conn.execute("SELECT COUNT(*) FROM van_ban").fetchone()[0] == 1


def test_lookup_exact_and_as_of(loaded):
    conn, _, _ = loaded
    (vb,) = find_van_ban(conn, "bltn", 2098)
    assert find_van_ban(conn, "BLTN", 1999) == [] and find_van_ban(conn, "XXXX") == []
    d = get_dieu(conn, vb["id"], "2", "2100-01-01")
    assert d.tieu_de == "Nguyên tắc"
    assert [(k.ky_hieu, [p.ky_hieu for p in k.children]) for k in d.children] == [("1", ["a", "b"]), ("2", [])]
    text = render(d)
    assert text.splitlines()[0] == "Điều 2. Nguyên tắc"
    assert "1. Nguyên tắc như sau:" in text and "a) tuân thủ pháp luật;" in text
    assert get_dieu(conn, vb["id"], "2a", "2100-01-01").noi_dung == "Nội dung 2a."
    assert get_dieu(conn, vb["id"], "1", "2100-01-01").tieu_de == "Phạm vi"          # Ð đã được sửa khi nạp
    assert get_dieu(conn, vb["id"], "9", "2100-01-01") is None                        # không có Điều 9
    assert get_dieu(conn, vb["id"], "2", "2099-06-30") is None                        # trước ngày hiệu lực
    assert get_dieu(conn, vb["id"], "2", "2099-07-01") is not None                    # ngày hiệu lực: đã có hiệu lực (nửa mở)


def test_lookup_ignores_versions_outside_date(loaded):
    conn, _, _ = loaded
    (vb,) = find_van_ban(conn, "BLTN")
    pb = conn.execute(
        "SELECT pb.* FROM phien_ban_noi_dung pb JOIN don_vi_luat dv ON dv.id = pb.don_vi_id "
        "WHERE dv.loai = 'dieu' AND pb.ky_hieu = '2a'").fetchone()
    with conn:
        conn.execute("UPDATE phien_ban_noi_dung SET hieu_luc_den = '2100-01-01' WHERE id = ?", (pb["id"],))
        conn.execute(
            "INSERT INTO phien_ban_noi_dung (don_vi_id, ky_hieu, thu_tu, tieu_de, noi_dung, hieu_luc_tu, van_ban_tao_ra_id) "
            "VALUES (?, '2a', ?, ?, 'Nội dung 2a mới', '2100-01-01', ?)",
            (pb["don_vi_id"], pb["thu_tu"], pb["tieu_de"], vb["id"]))
    assert get_dieu(conn, vb["id"], "2a", "2099-12-31").noi_dung == "Nội dung 2a."
    assert get_dieu(conn, vb["id"], "2a", "2100-01-01").noi_dung == "Nội dung 2a mới"
    assert get_dieu(conn, vb["id"], "2a") is None       # mặc định là hôm nay; luật thử nghiệm này hiệu lực năm 2099


@pytest.mark.parametrize("text,expected", [
    ("Điều 173 BLHS 2015", ("173", "BLHS", 2015)),
    ("điều 173a blhs", ("173a", "BLHS", None)),
    ("  Điều 5 BLDS 2015  ", ("5", "BLDS", 2015)),
    ("Điều BLHS", None),
    ("Khoản 2 Điều 5 BLHS", None),
    ("tội trộm cắp tài sản", None),
])
def test_parse_reference(text, expected):
    assert parse_reference(text) == expected


# ---------- toàn vẹn: không đánh rơi dòng nào ----------

def covered_lines(law):
    covered = set()
    for n in law.iter_nodes():
        covered.update(range(n.line_start, n.line_end + 1))
    return covered


def test_every_body_line_is_owned_by_a_node(staged):
    law = load_parsed(staged.parsed_path)
    body_end = len(law.lines) - len(law.footer)
    expected = set(range(len(law.preamble), body_end))
    assert covered_lines(law) == expected


# ---------- file BLHS 2015 thật (chỉ chạy khi có trong inbox/) ----------

@pytest.fixture(scope="module")
def real(tmp_path_factory):
    if not REAL_DOCX.exists():
        pytest.skip("không có inbox/Bộ-luật-100-2015-QH13.docx")
    tmp = tmp_path_factory.mktemp("real")
    st = stage_base_law(REAL_DOCX, tmp / "raw", ten_viet_tat="BLHS")
    errors = [f for f in st.flags if f.level == LOI]
    review = new_review(st.parsed_sha256)
    review["mien_loi"] = {flag_key(f): "kiểm thử: lỗi do văn bản nguồn" for f in errors}
    review["da_xem_canh_bao"] = [flag_key(f) for f in st.flags if f.level == CANH_BAO]
    save_review(approve(review, "pytest"), st.directory / "review.json")
    c = connect()
    init_db(c)
    ingest_base_law(c, st.directory, raw_root=tmp / "raw")
    yield st, c, errors
    c.close()


def test_real_counts_and_only_known_source_anomaly(real):
    st, c, errors = real
    counts = {}
    for n in st.law.iter_nodes():
        counts[n.loai] = counts.get(n.loai, 0) + 1
    assert counts["dieu"] == 426 and counts["phan"] == 3 and counts["chuong"] == 26
    assert (st.law.meta.so_hieu, st.law.meta.ngay_ban_hanh, st.law.meta.ngay_hieu_luc) == ("100/2015/QH13", "2015-11-27", "2016-07-01")
    assert st.law.source["homoglyphs_fixed"] >= 1
    # duy nhất một lỗi cấu trúc, do chính file nguồn (Điều 278 khoản 2 ghi e) ngay sau d))
    assert len(errors) == 1 and "Điều 278" in errors[0].message


def test_real_all_dieu_numbers_present_including_317(real):
    st, c, _ = real
    dieus = [n.ky_hieu for n in st.law.iter_nodes() if n.loai == "dieu"]
    assert dieus == [str(i) for i in range(1, 427)]
    (vb,) = find_van_ban(c, "BLHS", 2015)
    assert get_dieu(c, vb["id"], "317", "2020-01-01").tieu_de


def test_real_exact_lookup_works_for_every_dieu(real):
    st, c, _ = real
    (vb,) = find_van_ban(c, "BLHS")
    for i in range(1, 427):
        unit = get_dieu(c, vb["id"], str(i), "2020-01-01")
        assert unit is not None and unit.tieu_de, f"Điều {i}"
    assert get_dieu(c, vb["id"], "427", "2020-01-01") is None
    assert get_dieu(c, vb["id"], "1", "2016-06-30") is None       # trước hiệu lực THEO VĂN BẢN (chưa tính nghị quyết hoãn)


def test_real_every_body_line_owned(real):
    st, _, _ = real
    body_end = len(st.law.lines) - len(st.law.footer)
    assert covered_lines(st.law) == set(range(len(st.law.preamble), body_end))


def test_real_db_text_equals_source_for_every_dieu(real):
    """Parse -> SQLite -> render phải cho lại đúng văn bản nguồn của từng Điều (sau chuẩn hoá khoảng trắng)."""
    st, c, _ = real
    (vb,) = find_van_ban(c, "BLHS")
    law = st.law
    heads = sorted(
        (n for n in law.iter_nodes() if n.loai in ("phan", "chuong", "muc", "tieu_muc", "dieu")), key=lambda n: n.line_start
    )
    body_end = len(law.lines) - len(law.footer)
    bad = []
    for k, node in enumerate(heads):
        if node.loai != "dieu":
            continue
        end = heads[k + 1].line_start if k + 1 < len(heads) else body_end
        source = normalize_text(" ".join(law.lines[node.line_start:end]))
        stored = normalize_text(render(get_dieu(c, vb["id"], node.ky_hieu, "2020-01-01")))
        if source != stored:
            bad.append(node.ky_hieu)
    assert bad == []


def test_real_sample_for_human_spot_check_is_fixed(real):
    st, _, _ = real
    sample = sample_dieu(st.law)
    assert len(sample) == 30 and len(set(sample)) == 30
    assert sample == sample_dieu(st.law)                       # cùng seed, cùng danh sách


def test_real_warnings_are_the_four_split_points(real):
    st, _, _ = real
    warns = [f.message for f in st.flags if f.level == CANH_BAO]
    assert len(warns) == 4 and all("tách điểm" in w for w in warns)
