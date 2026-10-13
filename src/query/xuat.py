"""Xuất bản hợp nhất của một văn bản tại một ngày ra Markdown / Word (.docx): toàn bộ Phần/Chương/Mục/Điều/khoản/điểm đang có hiệu lực,
kèm NGUỒN sửa đổi từng Điều, phụ lục hiệu đính kỹ thuật và (nếu có) kết quả đối chiếu tự động với văn bản hợp nhất chính thức.

Bản xuất là dữ liệu dẫn xuất (dựng từ CSDL), KHÔNG phải văn bản hợp nhất chính thức; phần đầu bản xuất nói rõ điều này cùng các giới hạn.
Không dùng LLM. Chỉ đọc CSDL.
"""
from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from src.query.lookup import UnitView, fmt_date, get_cay, luu_y_hieu_luc, nhan_duyet, so_hieu_dinh_cho_duyet, so_lenh_cho_duyet

_TEN = {"phan": "Phần", "chuong": "Chương", "muc": "Mục", "tieu_muc": "Tiểu mục"}
_CAP = {"phan": 2, "chuong": 3, "muc": 4, "tieu_muc": 5, "dieu": 6}


@dataclass
class BoLuat:
    so_hieu: str
    ten_day_du: str
    ngay: str
    ngay_hieu_luc_goc: str | None
    roots: list[UnitView]
    cac_van_ban: list[dict] = field(default_factory=list)        # văn bản sửa đổi đã duyệt và có hiệu lực tại ngày: so_hieu, ten, ngay_ban_hanh, tu, den, n_lenh
    hieu_dinh: list[dict] = field(default_factory=list)          # quy tắc hiệu đính đã áp: id, khoa, tu_cu, thanh_moi, loai, nguon_so_sanh, nguoi_duyet, da_kiem_nguon
    n_dieu: int = 0
    n_dieu_da_sua: int = 0
    lenh_cho_duyet: int = 0
    hieu_dinh_cho_duyet: int = 0
    luu_y: str | None = None
    doi_chieu: dict | None = None                                # kết quả đối chiếu tự động: nhan, n_theirs, n_match, n_match_fold, n_noi_dung, n_trinh_bay, dieu_lech
    xem_truoc: bool = False                                      # bản chạy thử: GỒM cả lệnh/hiệu đính chưa được người duyệt


def _nhan_nut(path: list[UnitView]) -> str:
    """'khoản 2', 'điểm a khoản 1', 'toàn Điều' (path đi từ Điều xuống nút)."""
    if len(path) == 1:
        return "toàn Điều"
    return " ".join(("khoản" if u.loai == "khoan" else "điểm") + f" {u.ky_hieu}" for u in reversed(path[1:]))


def _ghi_chu_dieu(dieu: UnitView) -> tuple[list[tuple[str, str]], list[tuple[str, dict]]]:
    """(nguồn sửa đổi, hiệu đính) của Điều: [(nhãn nút, mô tả)], bỏ nút cùng nguồn với nút cha để khỏi lặp ở từng khoản/điểm con."""
    nguon: list[tuple[str, str]] = []
    hieu_dinh: list[tuple[str, dict]] = []

    def walk(u: UnitView, path: list[UnitView], parent_key: tuple[int, ...]) -> None:
        key = tuple(n.lenh_id for n in u.nguon)
        if u.nguon and key != parent_key:
            nguon.append((_nhan_nut(path), "; rồi ".join(n.mo_ta() for n in u.nguon)))
        for r in u.hieu_dinh:
            hieu_dinh.append((_nhan_nut(path), r))
        for c in u.children:
            walk(c, path + [c], key)

    walk(dieu, [dieu], ())
    return nguon, hieu_dinh


def _mo_ta_hieu_dinh(nhan: str, r: dict) -> str:
    base = (r.get("nguon_so_sanh") or "").split(" (")[0] or "văn bản hợp nhất"
    kiem = ", đã kiểm Công báo" if r.get("da_kiem_nguon") else ""
    return f"{nhan}: “{r.get('tu_cu', '?')}” → “{r.get('thanh_moi', '?')}” (theo {base}{kiem}; {nhan_duyet(r)})"


def dung_bo_luat(conn: sqlite3.Connection, van_ban_id: int, ngay: str | None = None, doi_chieu: dict | None = None,
                 xem_truoc: bool = False) -> BoLuat:
    """xem_truoc=True: `conn` là bản chạy thử (rebuild.preview) đã áp cả lệnh/hiệu đính chưa duyệt; bản xuất ghi rõ điều đó."""
    ngay = ngay or date.today().isoformat()
    vb = conn.execute("SELECT so_hieu, ten_day_du, ngay_hieu_luc FROM van_ban WHERE id = ?", (van_ban_id,)).fetchone()
    if vb is None:
        raise ValueError(f"Không có văn bản id {van_ban_id}")
    roots = get_cay(conn, van_ban_id, ngay)
    bl = BoLuat(vb["so_hieu"], vb["ten_day_du"], ngay, vb["ngay_hieu_luc"], roots, doi_chieu=doi_chieu, xem_truoc=xem_truoc)
    trang_thai = ("da_duyet", "da_sua", "cho_duyet") if xem_truoc else ("da_duyet", "da_sua")
    bl.cac_van_ban = [dict(r) for r in conn.execute(
        """SELECT v.so_hieu, v.ten_day_du AS ten, v.ngay_ban_hanh, MIN(l.hieu_luc_tu) AS tu, MAX(l.hieu_luc_tu) AS den, COUNT(*) AS n_lenh
           FROM lenh_sua_doi l JOIN van_ban v ON v.id = l.van_ban_nguon_id
           WHERE l.van_ban_dich_id = ? AND l.trang_thai_duyet IN (%s) AND l.hieu_luc_tu <= ?
           GROUP BY v.id ORDER BY tu, v.so_hieu""" % ",".join("?" * len(trang_thai)), (van_ban_id, *trang_thai, ngay))]
    seen: set[int] = set()

    def walk(u: UnitView) -> None:
        for c in u.children:
            walk(c)
    for root in roots:
        for u in root.walk():
            if u.loai == "dieu":
                bl.n_dieu += 1
                nguon, hd = _ghi_chu_dieu(u)
                bl.n_dieu_da_sua += bool(nguon)
                for nhan, r in hd:
                    if r.get("id") not in seen:
                        seen.add(r.get("id"))
                        bl.hieu_dinh.append({**r, "dieu": u.ky_hieu, "nhan": nhan})
    bl.hieu_dinh.sort(key=lambda r: r.get("id", 0))
    bl.lenh_cho_duyet = so_lenh_cho_duyet(conn, van_ban_id)
    bl.hieu_dinh_cho_duyet = so_hieu_dinh_cho_duyet(conn, van_ban_id)
    bl.luu_y = luu_y_hieu_luc(conn, van_ban_id)
    return bl


# ---------------------------------------------------------------------------------------------------------------------
# Khối nội dung trung gian (dùng chung cho Markdown và Word)
# ("tieu_de_van_ban", cấp, text) | ("tieu_de", cấp, text) | ("doan", thụt, text) | ("ghi_chu", 0, text) | ("muc_lon", cấp, text) | ("bang", 0, rows)


def _gioi_thieu(bl: BoLuat) -> list[tuple]:
    b: list[tuple] = []
    if bl.xem_truoc:
        b.append(("doan", 0, f"BẢN XEM TRƯỚC — CHƯA ĐƯỢC NGƯỜI DUYỆT. Bản này là kết quả chạy thử của hệ thống, đã gồm {bl.lenh_cho_duyet} lệnh sửa đổi và "
                             f"{bl.hieu_dinh_cho_duyet} quy tắc hiệu đính kỹ thuật chưa được người duyệt xác nhận. Chỉ dùng để xem và đối chiếu; "
                             "sau khi người duyệt (scripts/apply_commands.py duyet, scripts/doi_chieu_vbhn.py duyet) hãy xuất lại để có bản chính thức trong CSDL."))
    b.append(("doan", 0, f"Bản hợp nhất tại ngày {fmt_date(bl.ngay)} của {bl.ten_day_du} (số {bl.so_hieu}), do hệ thống Lexo dựng từ văn bản gốc và "
                         "các văn bản sửa đổi, bổ sung đã được người dùng duyệt. Đây là bản dẫn xuất phục vụ tra cứu, KHÔNG phải văn bản hợp nhất chính thức; "
                         "văn bản có giá trị là bản đăng Công báo và văn bản hợp nhất do Văn phòng Quốc hội xác thực."))
    b.append(("doan", 0, f"Quy mô: {bl.n_dieu} Điều đang có hiệu lực tại ngày này, trong đó {bl.n_dieu_da_sua} Điều đã được sửa đổi, bổ sung, bãi bỏ một phần "
                         "hoặc bổ sung mới sau khi ban hành (mỗi Điều ghi rõ nguồn)."))
    rows = [["Văn bản", "Ngày ban hành", "Hiệu lực từ", "Số lệnh áp dụng"]]
    rows.append([f"{bl.ten_day_du} (số {bl.so_hieu}) — văn bản gốc", "", fmt_date(bl.ngay_hieu_luc_goc) if bl.ngay_hieu_luc_goc else "?", ""])
    for v in bl.cac_van_ban:
        tu = fmt_date(v["tu"]) if v["tu"] == v["den"] else f"{fmt_date(v['tu'])} đến {fmt_date(v['den'])} (theo từng nội dung)"
        rows.append([f"{v['ten']} (số {v['so_hieu']})", fmt_date(v["ngay_ban_hanh"]) if v["ngay_ban_hanh"] else "?", tu, str(v["n_lenh"])])
    b.append(("muc_lon", 2, "Cơ sở hợp nhất"))
    b.append(("bang", 0, rows))
    d = bl.doi_chieu
    b.append(("muc_lon", 2, "Đối chiếu độc lập với văn bản hợp nhất chính thức"))
    if d:
        dong = [f"Đối chiếu từng Phần/Chương/Mục/Điều/khoản/điểm với {d['nhan']} (bản hệ thống tại ngày {fmt_date(d['ngay'])}): "
                f"{d['n_match_fold']}/{d['n_theirs']} nút khớp ({d['n_match_fold'] / max(d['n_theirs'], 1):.2%}), trong đó {d['n_match']} nút giống hệt từng chữ "
                f"và {d['n_match_fold'] - d['n_match']} nút chỉ khác cách trình bày (viết hoa, vị trí dấu thanh oá/óa, sỹ/sĩ, dấu câu, số 0 đầu số)."]
        if d["n_noi_dung"]:
            dong.append(f"Còn {d['n_noi_dung']} chỗ khác về nội dung ở các Điều: {', '.join(d['dieu_lech'][:40])}"
                        + (" …" if len(d["dieu_lech"]) > 40 else "") + " — chưa được giải quyết, cần người đối chiếu xem.")
        else:
            dong.append("Không còn chỗ nào khác về nội dung.")
        if bl.hieu_dinh:
            dong.append(f"Sự khớp này đạt được SAU {len(bl.hieu_dinh)} hiệu đính kỹ thuật (sửa chữ có bằng chứng, từng quy tắc do người duyệt, liệt kê ở Phụ lục); "
                        "nguyên văn như luật ban hành vẫn giữ trong CSDL (tra bằng scripts/lookup.py --nguyen-van).")
        b += [("doan", 0, x) for x in dong]
    else:
        b.append(("doan", 0, "Chưa chạy đối chiếu với văn bản hợp nhất chính thức khi xuất bản này (cần raw/<văn bản hợp nhất>/vbhn.json; xem scripts/doi_chieu_vbhn.py)."))
    b.append(("muc_lon", 2, "Giới hạn cần biết"))
    gh = ["Điều khoản hiệu lực và chuyển tiếp của các luật sửa đổi chỉ được lưu nguyên văn, hệ thống KHÔNG tự suy ra việc áp dụng hồi tố hay chuyển tiếp "
          "(xem scripts/lookup.py --chuyen-tiep)."]
    if bl.luu_y:
        y = bl.luu_y.removeprefix("Lưu ý: ")
        gh.append(y[:1].upper() + y[1:])
    if bl.lenh_cho_duyet and not bl.xem_truoc:
        gh.append(f"Còn {bl.lenh_cho_duyet} lệnh sửa đổi đã nạp nhưng CHƯA được duyệt, nên chưa nằm trong bản này.")
    if bl.hieu_dinh_cho_duyet and not bl.xem_truoc:
        gh.append(f"Còn {bl.hieu_dinh_cho_duyet} đề xuất hiệu đính kỹ thuật CHƯA được duyệt, nên chưa nằm trong bản này.")
    b += [("doan", 0, "• " + x) for x in gh]
    return b


def _doan(text: str, thut: int, thut_tiep: int) -> list[tuple]:
    """Văn bản nhiều đoạn (xuống dòng trong nội dung) -> mỗi đoạn một khối; đoạn đầu thụt `thut`, các đoạn sau `thut_tiep`."""
    parts = [x.strip() for x in text.split("\n") if x.strip()]
    return [("doan", thut if i == 0 else thut_tiep, x) for i, x in enumerate(parts)]


def _khoi_than(u: UnitView) -> list[tuple]:
    b: list[tuple] = []
    if u.loai in _TEN:
        ten = f"{_TEN[u.loai]} {u.ky_hieu}" + (f" — {u.tieu_de}" if u.tieu_de else "")
        b.append(("tieu_de", _CAP[u.loai], ten))
    elif u.loai == "dieu":
        b.append(("tieu_de", _CAP["dieu"], f"Điều {u.ky_hieu}. {u.tieu_de or ''}".rstrip()))
        if u.noi_dung:
            b += _doan(u.noi_dung, 0, 0)
        nguon, hd = _ghi_chu_dieu(u)
        # ghi chú đặt SAU nội dung Điều: thêm sau khi duyệt hết con
        ghi_chu = [("ghi_chu", 0, f"Nguồn sửa đổi — {nhan}: {mo_ta}") for nhan, mo_ta in nguon]
        ghi_chu += [("ghi_chu", 0, "Hiệu đính kỹ thuật — " + _mo_ta_hieu_dinh(nhan, r)) for nhan, r in hd]
        for c in u.children:
            b += _khoi_than(c)
        return b + ghi_chu
    elif u.loai == "khoan":
        b += _doan(f"{u.ky_hieu}. {u.noi_dung}".rstrip(), 0, 0)
    elif u.loai == "diem":
        b += _doan(f"{u.ky_hieu}) {u.noi_dung}".rstrip(), 1, 0)       # đoạn nối tiếp sau danh sách điểm thuộc cấp khoản, không thụt
    for c in u.children:
        b += _khoi_than(c)
    return b


def _phu_luc(bl: BoLuat) -> list[tuple]:
    b: list[tuple] = [("muc_lon", 2, "Phụ lục: hiệu đính kỹ thuật đã áp")]
    if not bl.hieu_dinh:
        b.append(("doan", 0, "Không có hiệu đính kỹ thuật nào: văn bản đúng như luật ban hành cộng các sửa đổi đã duyệt."))
        return b
    b.append(("doan", 0, "Mỗi dòng là một sửa chữ nhỏ so với văn bản NHƯ LUẬT BAN HÀNH (lỗi gõ trong bản đăng Công báo hoặc trong file Word luật sửa đổi, "
                         "hay con số chép sai), chỉ áp sau khi người duyệt xác nhận bằng chứng."))
    rows = [["#", "Điều / nút", "Từ cũ", "Thành", "Loại", "Căn cứ", "Duyệt"]]
    ten = {"loi_go": "lỗi gõ", "nhieu_chu": "nhiều chữ", "sai_so_lieu": "số liệu"}
    for r in bl.hieu_dinh:
        can_cu = (r.get("nguon_so_sanh") or "").split(" (")[0]
        if r.get("dong_thuan"):
            can_cu += "; đồng thuận: " + ", ".join(r["dong_thuan"])
        if r.get("da_kiem_nguon"):
            can_cu += "; đã kiểm nguồn: " + " | ".join(r["da_kiem_nguon"])
        rows.append([str(r.get("id")), f"Điều {r['dieu']} — {r['nhan']}", f"“{r.get('tu_cu')}”", f"“{r.get('thanh_moi')}”", ten.get(r.get("loai"), r.get("loai", "")),
                     can_cu, f"{r['nguoi_duyet']} ({fmt_date(r.get('ngay_duyet'))})" if r.get("nguoi_duyet") and r.get("ngay_duyet") else (r.get("nguoi_duyet") or "CHƯA DUYỆT")])
    b.append(("bang", 0, rows))
    return b


def khoi_noi_dung(bl: BoLuat) -> list[tuple]:
    out: list[tuple] = [("tieu_de_van_ban", 1, f"{bl.ten_day_du.upper()} — BẢN HỢP NHẤT TẠI NGÀY {fmt_date(bl.ngay)}" + (" (BẢN XEM TRƯỚC, CHƯA DUYỆT)" if bl.xem_truoc else ""))]
    out += _gioi_thieu(bl)
    out.append(("muc_lon", 2, "Nội dung"))
    for r in bl.roots:
        out += _khoi_than(r)
    out += _phu_luc(bl)
    return out


# ---------------------------------------------------------------------------------------------------------------------
# Markdown

def _md_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def ra_markdown(bl: BoLuat) -> str:
    lines: list[str] = []
    for kind, n, val in khoi_noi_dung(bl):
        if kind in ("tieu_de_van_ban", "muc_lon"):
            lines += ["#" * n + " " + val, ""]
        elif kind == "tieu_de":
            lines += ["#" * n + " " + val, ""]
        elif kind == "doan":
            text = val
            head, sep, rest = text.partition(". ")
            if head.isdigit() and sep:
                text = f"{head}\\. {rest}"                           # tránh bị Markdown đổi "1. ..." thành danh sách đánh số lại
            lines += [("&emsp;" * 2 * n) + text if n else text, ""]
        elif kind == "ghi_chu":
            lines += [f"> *{val}*", ""]
        elif kind == "bang":
            rows = val
            lines.append("| " + " | ".join(_md_cell(c) for c in rows[0]) + " |")
            lines.append("|" + "|".join(["---"] * len(rows[0])) + "|")
            lines += ["| " + " | ".join(_md_cell(c) for c in r) + " |" for r in rows[1:]]
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------------------------------------------------
# Word

def ra_docx(bl: BoLuat, path: str | Path) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for s in d.sections:
        s.left_margin, s.right_margin = Cm(2.5), Cm(2)
    for kind, n, val in khoi_noi_dung(bl):
        if kind == "tieu_de_van_ban":
            h = d.add_heading(val, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind in ("muc_lon", "tieu_de"):
            level = min(max(n - 1, 1), 5) if kind == "tieu_de" else 1
            d.add_heading(val, level=level)
        elif kind == "doan":
            p = d.add_paragraph(val)
            p.paragraph_format.space_after = Pt(4)
            if n:
                p.paragraph_format.left_indent = Cm(1.0 * n)
        elif kind == "ghi_chu":
            p = d.add_paragraph()
            run = p.add_run(val)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
        elif kind == "bang":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for i, row in enumerate(val):
                for j, cell in enumerate(row):
                    c = t.cell(i, j)
                    c.text = cell
                    for para in c.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
                            run.bold = i == 0
            d.add_paragraph()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path
