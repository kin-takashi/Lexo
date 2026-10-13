"""Trích text thuần từ file Word (.docx): danh sách dòng đã chuẩn hoá + nội dung bảng (số hiệu, quốc hiệu...)."""
from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

from src.parse.text import count_homoglyphs, normalize_text


@dataclass
class ExtractedText:
    lines: list[str]           # đoạn không rỗng, đã chuẩn hoá, theo thứ tự trong văn bản
    header_text: str           # nội dung các bảng ghép lại (số hiệu văn bản nằm ở đây)
    source_name: str
    sha256: str                # của file gốc, để chắc file được duyệt là file được lưu vào raw/
    homoglyphs_fixed: int      # số ký tự Ð/ð đã sửa thành Đ/đ
    n_tables: int

    def source_info(self) -> dict:
        return {"name": self.source_name, "sha256": self.sha256,
                "homoglyphs_fixed": self.homoglyphs_fixed, "n_tables": self.n_tables}


def extract_docx(path: str | Path) -> ExtractedText:
    import docx  # python-docx

    path = Path(path)
    data = path.read_bytes()
    document = docx.Document(io.BytesIO(data))

    raw_lines = [p.text for p in document.paragraphs]
    table_texts: list[str] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                if text.strip() and (not table_texts or table_texts[-1] != text):   # ô gộp lặp lại cùng nội dung
                    table_texts.append(text)

    fixed = sum(count_homoglyphs(t) for t in raw_lines) + sum(count_homoglyphs(t) for t in table_texts)
    lines = [n for n in (normalize_text(t) for t in raw_lines) if n]
    header_text = " ".join(n for n in (normalize_text(t) for t in table_texts) if n)
    return ExtractedText(
        lines=lines,
        header_text=header_text,
        source_name=path.name,
        sha256=hashlib.sha256(data).hexdigest(),
        homoglyphs_fixed=fixed,
        n_tables=len(document.tables),
    )
